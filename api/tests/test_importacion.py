# Copyright (c) 2026 Jose Danilo Camacho A. / Tecno-Insights S.A.S.
# SPDX-License-Identifier: MIT

"""Test de `GET /api/importacion/excel/plantilla` (ADR-0008 F3, P4).

`exportar_plantilla` movió la construcción del libro de Excel (openpyxl,
síncrono) a un hilo (`asyncio.to_thread` sobre `_construir_libro_plantilla`);
este test confirma que la respuesta sigue siendo un .xlsx válido con los datos
del requerimiento.
"""
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.documents.requerimiento import Requerimiento, Solicitud
from tests.conftest import headers_con_token


async def test_exportar_plantilla_devuelve_un_xlsx_valido_con_los_datos(
    cliente, fabrica_usuario, fabrica_aplicacion
):
    app = await fabrica_aplicacion()
    await Requerimiento(
        aplicacion_id=app.codigo,
        codigo_req="REQ-PLANTILLA-1",
        solicitud=Solicitud(codigo_sc="SC-PLANTILLA-1"),
        estado="ABIERTO",
        total_horas_estimadas=Decimal("12"),
    ).insert()
    _, token = await fabrica_usuario([app.codigo], permisos=["admin.importacion.ver"])

    resp = await cliente.get(
        "/api/importacion/excel/plantilla", headers=headers_con_token(token, app.codigo)
    )
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    libro = load_workbook(BytesIO(resp.content))
    assert {"REQUERIMIENTOS", "ENTREGAS"} <= set(libro.sheetnames)
    hoja = libro["REQUERIMIENTOS"]
    filas = list(hoja.iter_rows(min_row=2, values_only=True))
    assert any(fila[1] == "REQ-PLANTILLA-1" for fila in filas)
